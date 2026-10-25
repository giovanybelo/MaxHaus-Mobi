# -*- coding: utf-8 -*-
"""Gera Orcamento_Banheiro_R2.xlsx a partir do R1, aplicando a paginação R1 do caderno."""
import sys, copy, openpyxl
from openpyxl.utils import get_column_letter
sys.path.insert(0, __import__('os').path.dirname(__import__('os').path.abspath(__file__)))
import calc_r1 as K

SRC, DST = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC)
HDR = 'MAXHAUS | R2 | 25/10/2026 | valores em R$ | conferir em obra'
for ws in wb.worksheets:
    if isinstance(ws['A2'].value, str) and 'MAXHAUS' in ws['A2'].value:
        ws['A2'] = HDR

def like(ws, src, dst):
    s = ws[src]; d = ws[dst]
    d.font = copy.copy(s.font); d.fill = copy.copy(s.fill); d.border = copy.copy(s.border)
    d.alignment = copy.copy(s.alignment); d.number_format = s.number_format

rows, res = K.totals()
R = {r['id']: r for r in rows}

# ---------------------------------------------------------------- Paredes
ws = wb['Paredes']
ws['E4'] = 'Área revestida m²'; ws['F4'] = 'Área Bleu m²'; ws['G4'] = 'Área Nuage m²'
for col, title in (('J', 'Peças inteiras'), ('K', 'Peças cortadas'), ('L', 'Corte das pontas mm')):
    ws[col + '4'] = title; like(ws, 'I4', col + '4')
    ws.column_dimensions[col].width = max(ws.column_dimensions['I'].width or 12, 14)
order = ['P1', 'P2', 'P3', 'P4', 'P5', 'P6', 'N1', 'N2', 'N3', 'N4', 'N5', 'B1', 'B2', 'B3', 'B4']
for i, sid in enumerate(order):
    r = R[sid]; row = 5 + i
    ws.cell(row, 3).value = r['W']; ws.cell(row, 4).value = r['H']
    ws.cell(row, 5).value = round(r['area'], 4)
    ws.cell(row, 6).value = round(r['area'] - r['area_n'], 4)
    ws.cell(row, 7).value = round(r['area_n'], 4)
    ws.cell(row, 8).value = r['nb']; ws.cell(row, 9).value = r['nn']
    ws.cell(row, 10).value = r['full']; ws.cell(row, 11).value = r['cut']
    ws.cell(row, 12).value = (round(r['c'], 1) if r['c'] is not None else '—')
    for col in ('J', 'K', 'L'):
        like(ws, 'I%d' % row, '%s%d' % (col, row))
ws['A21'] = ('R2: paredes até 2.353 mm (15 fiadas inteiras), pontas simétricas no eixo, nicho e porta preservados. '
             'Nicho mantém 1000 × 400 × 118 mm. Área Nuage = acentos (peça 75 × 155) + bancada.')

# ---------------------------------------------------------------- Quantidades
ws = wb['Quantidades']
heads = ['Material / parâmetro', 'Peças físicas', 'Área revestida m²', 'Reserva', 'm² compra', 'm²/cx', 'Caixas', 'Peças/cx', 'Peças compradas', 'Folga sobre peças']
for k, h in enumerate(heads):
    c = ws.cell(4, 1 + k); c.value = h
    if k >= 7: like(ws, 'G4', c.coordinate)
for row, colH, colF in ((5, 'H', 'F'), (6, 'I', 'G')):
    ws['B%d' % row] = '=SUM(Paredes!%s5:%s19)' % (colH, colH)
    ws['C%d' % row] = '=SUM(Paredes!%s5:%s19)' % (colF, colF)
    ws['D%d' % row] = 0.1
    ws['E%d' % row] = '=C%d*(1+D%d)' % (row, row)
    ws['F%d' % row] = 0.3
    ws['G%d' % row] = '=ROUNDUP(E%d/F%d,0)' % (row, row)
    ws['H%d' % row] = 26
    ws['I%d' % row] = '=G%d*H%d' % (row, row)
    ws['J%d' % row] = '=I%d/B%d-1' % (row, row)
    for col, srcc in (('B', 'B'), ('C', 'G'), ('D', 'C'), ('E', 'G'), ('F', 'F'), ('G', 'E'), ('H', 'D'), ('I', 'B'), ('J', 'C')):
        like(ws, '%s%d' % (srcc, row) if srcc != col else '%s%d' % (col, row), '%s%d' % (col, row))
    ws['C%d' % row].number_format = '0.00'; ws['E%d' % row].number_format = '0.00'
    ws['G%d' % row].number_format = '0'; ws['I%d' % row].number_format = '0'; ws['J%d' % row].number_format = '0.0%'
    ws['D%d' % row].number_format = '0%'
for col in 'HIJ':
    ws.column_dimensions[col].width = 16
ws['A12'] = 'Junta nominal mm'; ws['B12'] = 2
ws['A14'] = 'Altura revestida mm (15 fiadas inteiras)'; ws['B14'] = 2353
ws['A15'] = 'Compra = área revestida × (1 + reserva), arredondada para caixa fechada. Conferida contra as peças físicas: cada peça cortada consome uma peça.'
ws['A16'] = 'Calibre real do lote difere do nominal (Nuage 76,5 × 155,5): refazer a paginação com a peça medida antes de cortar.'

# ---------------------------------------------------------------- Orcamento
ws = wb['Orcamento']
ws['D5'] = '=Quantidades!G5'; ws['D6'] = '=Quantidades!G6'
ws['H5'] = ('Área revestida %s m² + 10%% = %s m², caixa de 0,30 m² / 26 peças. Cobre as %d peças físicas da paginação R2 com %d%% de folga.' %
            (('%.2f' % res['B']['area']).replace('.', ','), ('%.2f' % res['B']['m2_compra']).replace('.', ','), res['B']['pcs_fisicas'], round(res['B']['folga'] * 100)))
ws['H6'] = ('Área %s m² (acentos + bancada) + 10%% = %s m². Cobre as %d peças físicas com %d%% de folga. Preço provisório igual ao Bleu; cotar.' %
            (('%.2f' % res['N']['area']).replace('.', ','), ('%.2f' % res['N']['m2_compra']).replace('.', ','), res['N']['pcs_fisicas'], round(res['N']['folga'] * 100)))
area_tot = sum(r['area'] for r in rows)
ws['H9'] = 'Reserva 5 kg/m² × %s m² = %s kg; definir produto com Portobello e impermeabilização.' % (
    ('%.2f' % area_tot).replace('.', ','), ('%.0f' % (5 * area_tot)))
ws['B12'] = 'Perfis de acabamento: topo do revestimento a 2.353 mm e bordas/cantos'
ws['D12'] = 13
ws['H12'] = 'Topo: 8,02 m (perímetro das paredes P1–P6) + 4 m de bordas e cantos; perfil reto inox/alumínio para borda de peça de 9 mm. Cantos a 45° dependem de ensaio de corte.'

# ---------------------------------------------------------------- Orcamento: argamassa e impermeabilização
ws = wb['Orcamento']
ws['D9'] = 0; ws['G9'] = 'Contratado'
ws['H9'] = 'Fornecida pela empreiteira (contrato fechado). Referência de consumo: 5 kg/m² × %s m² ≈ %s kg; produto compatível com a impermeabilização.' % (
    ('%.2f' % area_tot).replace('.', ','), ('%.0f' % (5 * area_tot)))
IMP = [
    ('Impermeabilização', 'Impermeabilizante cimentício flexível, balde 18 kg (ref. Quartzolit Tecplus Lastic)', 'bd', 3, 199.90, 'Preço web',
     '12 m² (piso, box até 2.400 mm, rodapé técnico, bancada e nicho) × cerca de 3,5 kg/m² + 10% ≈ 46 kg. Sistema, demãos e consumo a validar com fabricante e aplicador.',
     'https://www.telhanorte.com.br/impermeabilizante-quartzolit-18kg-cinza-tecplus-lastic-1619934/p'),
    ('Impermeabilização', 'Tela estruturante de poliéster 73 cm × 10 m', 'rl', 1, 67.90, 'Preço web',
     'Reforço em cantos, encontro piso/parede, ralos R1 e R2, flanges e passagens (cerca de 25 m de faixa).',
     'https://www.telhanorte.com.br/tela-estruturante-quartzolit-1751182/p?skuId=1751182'),
]
for i, (g, it, un, q, pu, org, crit, fonte) in enumerate(IMP):
    r = 41 + i
    vals = [g, it, un, q, pu, '=ROUND(D%d*E%d,2)' % (r, r), org, crit, fonte]
    for k, v in enumerate(vals):
        cell = ws.cell(r, 1 + k); cell.value = v
        src = ws.cell(40, 1 + k)
        cell.font = copy.copy(src.font); cell.fill = copy.copy(src.fill); cell.border = copy.copy(src.border)
        cell.alignment = copy.copy(src.alignment); cell.number_format = src.number_format

# ---------------------------------------------------------------- Resumo (nova categoria)
ws = wb['Resumo']
st = {r: [copy.copy(ws.cell(r, k)._style) for k in (1, 2)] for r in range(5, 13)}
cats = ['Revestimentos', 'Louças e metais', 'Ralos e esgoto', 'Água quente/fria', 'Impermeabilização']
for i, cat in enumerate(cats):
    r = 5 + i
    ws.cell(r, 1).value = cat
    ws.cell(r, 2).value = '=SUMIF(Orcamento!A5:A60,A%d,Orcamento!F5:F60)' % r
    ws.cell(r, 1)._style = copy.copy(st[8][0]); ws.cell(r, 2)._style = copy.copy(st[8][1])
lines = [(10, 'Subtotal', '=SUM(B5:B9)', 9), (11, 'Contingência', 0.15, 10), (12, 'Reserva financeira R$', '=B10*B11', 11), (13, 'Total com contingência', '=B10+B12', 12)]
for r, a, b, src in lines:
    ws.cell(r, 1).value = a; ws.cell(r, 2).value = b
    ws.cell(r, 1)._style = copy.copy(st[src][0]); ws.cell(r, 2)._style = copy.copy(st[src][1])
ws['A14'] = ('Excluídos: piso Monofloor Everest, elétrica/iluminação, mão de obra, argamassa colante (empreiteira), '
             'logística (frete, entrega e subida no condomínio, caçamba de entulho) e pintura da faixa de 38 mm acima do revestimento.')
ws['A15'] = 'Reserva de 15% apresentada à parte; não é custo contratado.'
ws.row_dimensions[12].height = None; ws.row_dimensions[13].height = 32
ws.row_dimensions[14].height = 64; ws.row_dimensions[20].height = 48
for r_ in (14, 20):
    al = copy.copy(ws.cell(14, 1).alignment); al.wrap_text = True; al.vertical = 'top'; ws.cell(r_, 1).alignment = al
ws['A20'] = 'R2: paginação refeita (peças inteiras, topo a 2.353 mm, pontas simétricas), compra por área + 10%, argamassa com a empreiteira, impermeabilização incluída e sobreposições resolvidas.'

# ---------------------------------------------------------------- Premissas
ws = wb['Premissas']
ws['B6'] = 'R47; paredes Bleu Nuit com detalhes Nuage e tampo Nuage 7,5 × 15,5 cm. Paginação R2 no caderno técnico R1.'
ws['B19'] = 'Área revestida × 1,10, arredondada para caixa fechada; conferida contra as peças físicas da paginação. Compatibilizar bases, válvula/sifão, fixações e garantia antes da compra.'
for row, a, b in ((20, 'Paginação R2', '15 fiadas inteiras desde o piso; revestimento termina em 2.353 mm com perfil, faixa até o teto em pintura. Colunas simétricas no eixo com o maior número de peças inteiras e pontas ≥ 40% da peça. Nicho e porta preservados.'),
                  (21, 'Alternativa', 'Revestir até 2.391 mm exige fiada de 36 mm em todas as paredes (107 tiras, cerca de 27 peças) e +1 caixa de Bleu Nuit.')):
    ws['A%d' % row] = a; ws['B%d' % row] = b
    like(ws, 'A19', 'A%d' % row); like(ws, 'B19', 'B%d' % row)

# ---------------------------------------------------------------- Comparacao PB
ws = wb['Comparacao PB']
ws['A5'] = 'Bleu Nuit: oficial 16,20 m² / R2 %s m²' % ('%.2f' % (res['B']['caixas'] * 0.3)).replace('.', ',')
ws['A6'] = 'Nuage: não cotado / R2 %s m²' % ('%.2f' % (res['N']['caixas'] * 0.3)).replace('.', ',')
ws['A13'] = 'Diferença no Bleu: %d caixas a mais que a proposta oficial (%s m²), não aumento do preço por m².' % (
    res['B']['caixas'] - 54, ('%.2f' % ((res['B']['caixas'] - 54) * 0.3)).replace('.', ','))
ws['A8'] = 'Espaçadores, selante e perfis (argamassa: empreiteira)'
ws['C4'] = 'R2 R$'
ws['A14'] = '%d caixas = área revestida %s m² + 10%%, por caixa fechada. R1 previa 70 caixas (peça cortada + 10%%).' % (
    res['B']['caixas'], ('%.2f' % res['B']['area']).replace('.', ','))

# ---------------------------------------------------------------- nova aba: Revisao R2
ws = wb.create_sheet('Revisao R2')
src = wb['Comparacao PB']
ws['A1'] = 'O que mudou do R1 para o R2'; like_src = lambda a, b: None
ws['A1'].font = copy.copy(src['A1'].font)
ws['A2'] = HDR; ws['A2'].font = copy.copy(src['A2'].font)
heads = ['Item', 'R1', 'R2', 'Diferença', 'Por quê']
for k, h in enumerate(heads):
    c = ws.cell(4, 1 + k); c.value = h
    s = src.cell(4, min(1 + k, 4))
    c.font = copy.copy(s.font); c.fill = copy.copy(s.fill); c.border = copy.copy(s.border); c.alignment = copy.copy(s.alignment)
data = [
    ('Bleu Nuit, caixas', 70, '=Orcamento!D5', 'Compra por área + 10% e paginação com mais peças inteiras'),
    ('Nuage, caixas', 7, '=Orcamento!D6', 'Mesmos 152 acentos e bancada'),
    ('Peças físicas Bleu', 1632, '=Quantidades!B5', 'Sem fiada de 36 mm no topo e com pontas simétricas'),
    ('Peças físicas Nuage', 152, '=Quantidades!B6', ''),
    ('Perfis de acabamento, m', 4, '=Orcamento!D12', 'Perfil no topo do revestimento (8,02 m)'),
    ('Argamassa colante, R$', 325.00, '=Orcamento!F9', 'Passa para a empreiteira (já contratada)'),
    ('Revestimentos, R$', 9978.63, '=Resumo!B5', ''),
    ('Impermeabilização, R$', 0, '=Resumo!B9', 'Incluída no R2: 3 baldes de 18 kg e 1 rolo de tela'),
    ('Subtotal, R$', 20023.85, '=Resumo!B10', ''),
    ('Total com contingência, R$', 23027.43, '=Resumo!B13', 'Contingência de 15% mantida'),
]
for i, (a, b, c, d) in enumerate(data):
    r = 5 + i
    ws.cell(r, 1).value = a; ws.cell(r, 2).value = b; ws.cell(r, 3).value = c
    ws.cell(r, 4).value = '=C%d-B%d' % (r, r); ws.cell(r, 5).value = d
    for k in range(1, 6):
        s = src.cell(5, min(k, 4)); t = ws.cell(r, k)
        t.font = copy.copy(s.font); t.border = copy.copy(s.border); t.alignment = copy.copy(s.alignment)
    fmt = '#,##0.00' if 'R$' in a else '0'
    for k in (2, 3, 4): ws.cell(r, k).number_format = fmt
ws.column_dimensions['A'].width = 30; ws.column_dimensions['E'].width = 60
for col in 'BCD': ws.column_dimensions[col].width = 14
ws = wb.create_sheet('Sobreposicoes')
ws['A1'] = 'Itens que se sobrepõem: decisão do R2'; ws['A1'].font = copy.copy(src['A1'].font)
ws['A2'] = HDR; ws['A2'].font = copy.copy(src['A2'].font)
heads = ['Função', 'Opção mantida', 'R$ mantido', 'Opção descartada', 'R$ descartado', 'Por quê']
for k, h in enumerate(heads):
    c = ws.cell(4, 1 + k); c.value = h; s_ = src.cell(4, min(1 + k, 4))
    c.font = copy.copy(s_.font); c.fill = copy.copy(s_.fill); c.border = copy.copy(s_.border); c.alignment = copy.copy(s_.alignment)
SOB = [
    ('Ralo linear do box', 'Ralo linear 720 mm completo', '=Orcamento!F24', 'Ralo linear 600 × 90 mm do link', 199.00,
     'Projeto pede 720 mm. O de 600 mm não atende à medida e tem menor valor: descartado.'),
    ('Misturador da bancada', 'Bica 200 mm + 2 registros Round + bases 3 furos', '=Orcamento!F14+Orcamento!F15+Orcamento!F16', 'Royale monocomando 21 cm', 527.00,
     'O modelo desenha bica e dois comandos (3 furos). Royale é monocomando de 2 furos: não atende e tem menor valor.'),
    ('Ligação do vaso', 'Engate flexível certificado', '=Orcamento!F28', 'Tubo de ligação 1½"', 272.87,
     'Vaso com caixa acoplada: o tubo de 1½" é para válvula de descarga e não atende. Mantido o engate, mesmo com valor menor.'),
    ('Fixação da torneira', 'Fixação do gabarito do fabricante', 0, 'Rosca longa 12 cm', 36.75,
     'Só comprar se o gabarito real exigir. Fora do orçamento.'),
    ('Argamassa colante', 'Empreiteira (contratada)', '=Orcamento!F9', 'Argamassa no orçamento de materiais', 325.00,
     'Fornecida pela empreiteira: não somar duas vezes.'),
    ('Toalheiros', 'Aiello térmico + Bracci de rosto', '=Orcamento!F22+Orcamento!F23', '—', 0,
     'Não é sobreposição: o térmico é para toalha de banho junto ao box; o de rosto fica junto à bancada. Se quiser um só, sai o Bracci (−R$ 433,65).'),
]
for i, row in enumerate(SOB):
    r = 5 + i
    for k, v in enumerate(row):
        c = ws.cell(r, 1 + k); c.value = v; s_ = src.cell(5, min(1 + k, 4))
        c.font = copy.copy(s_.font); c.border = copy.copy(s_.border); c.alignment = copy.copy(s_.alignment)
    ws.cell(r, 3).number_format = '#,##0.00'; ws.cell(r, 5).number_format = '#,##0.00'
for col, wd in zip('ABCDEF', (24, 40, 13, 32, 13, 70)):
    ws.column_dimensions[col].width = wd
wb.save(DST)
print('ok', DST)
